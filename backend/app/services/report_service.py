"""
Report generation service — PDF executivo e Excel de alertas.
Suporta branding configurável via BRAND_* settings (white-label adriner.fr).
"""
from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from sqlalchemy.orm import Session, joinedload

from app.core.config import settings
from app.models.fiscal_document import FiscalDocument
from app.models.rule_log import RuleExecutionLog
from app.services.scoring_service import AlertSeverity, ScoringService


# ── Helpers ──────────────────────────────────────────────────────────────────

SEVERITY_ORDER = {
    AlertSeverity.CRITICAL: 0,
    AlertSeverity.HIGH: 1,
    AlertSeverity.MEDIUM: 2,
    AlertSeverity.LOW: 3,
    AlertSeverity.INFORMATIVE: 4,
}

SEVERITY_PT = {
    "CRITICAL": "Crítico",
    "HIGH": "Alto",
    "MEDIUM": "Médio",
    "LOW": "Baixo",
    "INFORMATIVE": "Informativo",
}


def _currency(value: float | Decimal) -> str:
    return f"R$ {float(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _get_period_alerts(db: Session, tenant_id: str, year: int, month: int) -> list[dict]:
    """Return all failed-rule alerts for the period, sorted by severity × exposure."""
    start = date(year, month, 1)
    end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)

    docs = (
        db.query(FiscalDocument)
        .options(joinedload(FiscalDocument.rule_logs))
        .filter(
            FiscalDocument.tenant_id == tenant_id,
            FiscalDocument.data_emissao >= start,
            FiscalDocument.data_emissao < end,
        )
        .all()
    )

    alerts: list[dict] = []
    for doc in docs:
        for log in doc.rule_logs:
            if log.passed:
                continue
            severity, exposure = ScoringService.calculate_alert_severity(log, doc)
            alerts.append(
                {
                    "rule_id": log.rule_id,
                    "severity": severity,
                    "severity_label": SEVERITY_PT.get(severity.value, severity.value),
                    "emitente": doc.emitente_nome or "–",
                    "chave_acesso": doc.chave_acesso or "–",
                    "data_emissao": doc.data_emissao.strftime("%d/%m/%Y") if doc.data_emissao else "–",
                    "valor_nf": float(doc.valor_total or 0),
                    "exposure": float(exposure),
                    "message": log.message or "",
                }
            )

    alerts.sort(key=lambda a: (SEVERITY_ORDER[a["severity"]], -a["exposure"]))
    return alerts


def _get_period_score(db: Session, tenant_id: str, year: int, month: int) -> dict:
    return ScoringService.get_period_score(db, tenant_id, year, month)


# ── PDF ───────────────────────────────────────────────────────────────────────

def generate_pdf_report(db: Session, tenant_id: str, year: int, month: int) -> bytes:
    """
    Gera relatório executivo em PDF para o período informado.
    Inclui branding configurável (BRAND_* settings) — ideal para adriner.fr white-label.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT

    period_data = _get_period_score(db, tenant_id, year, month)
    alerts = _get_period_alerts(db, tenant_id, year, month)

    score = period_data.get("period_score", 100)
    total_exposure = period_data.get("total_exposure", 0.0)
    docs_count = period_data.get("documents_processed", 0)
    critical_count = period_data.get("critical_documents", 0)
    top_rules = period_data.get("top_3_rules", [])

    brand = settings.BRAND_NAME or "FiscalAI"
    expert_name = settings.BRAND_EXPERT_NAME
    expert_title = settings.BRAND_EXPERT_TITLE
    whatsapp = settings.BRAND_WHATSAPP

    month_name = [
        "", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
        "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
    ][month]

    # Score color
    if score >= 80:
        score_color = colors.HexColor("#16a34a")
    elif score >= 50:
        score_color = colors.HexColor("#d97706")
    else:
        score_color = colors.HexColor("#dc2626")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    brand_color = colors.HexColor("#1e3a5f")

    title_style = ParagraphStyle("Title", parent=styles["Title"], textColor=brand_color, fontSize=22, spaceAfter=4)
    subtitle_style = ParagraphStyle("Sub", parent=styles["Normal"], textColor=colors.HexColor("#64748b"), fontSize=11, spaceAfter=2)
    section_style = ParagraphStyle("Section", parent=styles["Heading2"], textColor=brand_color, fontSize=13, spaceBefore=16, spaceAfter=6)
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, spaceAfter=4)
    small_style = ParagraphStyle("Small", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#64748b"))
    center_style = ParagraphStyle("Center", parent=styles["Normal"], alignment=TA_CENTER, fontSize=10)

    story = []

    # ── Header ──
    story.append(Paragraph(brand, title_style))
    story.append(Paragraph(f"Relatório de Auditoria Fiscal — {month_name}/{year}", subtitle_style))
    story.append(Paragraph(f"Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}", small_style))
    story.append(HRFlowable(width="100%", thickness=2, color=brand_color, spaceAfter=12))

    # ── Score KPIs ──
    story.append(Paragraph("Resumo do Período", section_style))

    score_label = "Excelente" if score >= 80 else ("Atenção" if score >= 50 else "Crítico")
    kpi_data = [
        ["Score de Risco", "Exposição Estimada", "Documentos", "Críticos"],
        [
            Paragraph(f'<font color="{score_color.hexval()}" size="20"><b>{score}/100</b></font><br/><font size="9">{score_label}</font>', center_style),
            Paragraph(f'<b>{_currency(total_exposure)}</b>', center_style),
            Paragraph(f'<b>{docs_count}</b>', center_style),
            Paragraph(f'<font color="#dc2626"><b>{critical_count}</b></font>', center_style),
        ],
    ]
    kpi_table = Table(kpi_data, colWidths=[4.2 * cm, 4.5 * cm, 3.5 * cm, 3.5 * cm])
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), brand_color),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#f8fafc"), colors.white]),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 0.5 * cm))

    # ── Top regras ──
    if top_rules:
        story.append(Paragraph("Regras com Mais Ocorrências", section_style))
        rule_data = [["Regra", "Ocorrências"]]
        for r in top_rules:
            rule_data.append([r["rule_id"], str(r["failures"])])
        rule_table = Table(rule_data, colWidths=[12 * cm, 3.7 * cm])
        rule_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), brand_color),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#f8fafc"), colors.white]),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(rule_table)
        story.append(Spacer(1, 0.5 * cm))

    # ── Alertas detalhados (top 20) ──
    top_alerts = alerts[:20]
    if top_alerts:
        story.append(Paragraph(f"Top Inconsistências Detectadas ({len(top_alerts)} de {len(alerts)})", section_style))
        alert_data = [["Severidade", "Regra", "Emitente", "Data", "Valor NF", "Exposição"]]
        for a in top_alerts:
            alert_data.append([
                a["severity_label"],
                a["rule_id"],
                a["emitente"][:28],
                a["data_emissao"],
                _currency(a["valor_nf"]),
                _currency(a["exposure"]),
            ])
        alert_table = Table(
            alert_data,
            colWidths=[2.2 * cm, 3.8 * cm, 4.5 * cm, 2.2 * cm, 2.5 * cm, 2.5 * cm],
        )

        def _row_bg(i: int, sev: str) -> colors.Color:
            if sev == "Crítico":
                return colors.HexColor("#fef2f2")
            if sev == "Alto":
                return colors.HexColor("#fff7ed")
            return colors.HexColor("#f8fafc") if i % 2 == 0 else colors.white

        row_styles = []
        for i, row in enumerate(top_alerts, start=1):
            bg = _row_bg(i, row["severity_label"])
            row_styles.append(("BACKGROUND", (0, i), (-1, i), bg))

        alert_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), brand_color),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            *row_styles,
        ]))
        story.append(alert_table)

    story.append(Spacer(1, cm))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#e2e8f0"), spaceAfter=8))

    # ── Footer com assinatura do especialista ──
    if expert_name:
        footer_lines = [f"<b>{expert_name}</b>"]
        if expert_title:
            footer_lines.append(expert_title)
        if whatsapp:
            footer_lines.append(f"WhatsApp: {whatsapp}")
        footer_lines.append(f"Relatório gerado por {brand} em {datetime.now().strftime('%d/%m/%Y')}")
        story.append(Paragraph("<br/>".join(footer_lines), small_style))
    else:
        story.append(Paragraph(f"Relatório gerado por {brand} em {datetime.now().strftime('%d/%m/%Y')}", small_style))

    if whatsapp:
        story.append(Spacer(1, 0.3 * cm))
        cta = f'Dúvidas sobre estes alertas? Fale com o especialista via WhatsApp: <b>{whatsapp}</b>'
        story.append(Paragraph(cta, ParagraphStyle("CTA", parent=body_style, textColor=brand_color)))

    doc.build(story)
    return buf.getvalue()


# ── Excel ─────────────────────────────────────────────────────────────────────

def generate_excel_report(db: Session, tenant_id: str, year: int, month: int) -> bytes:
    """Gera planilha Excel com todas as inconsistências do período."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    alerts = _get_period_alerts(db, tenant_id, year, month)
    period_data = _get_period_score(db, tenant_id, year, month)

    brand = settings.BRAND_NAME or "FiscalAI"
    month_name = [
        "", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
        "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
    ][month]

    wb = openpyxl.Workbook()

    # ── Aba 1: Resumo ──
    ws_sum = wb.active
    ws_sum.title = "Resumo"

    header_fill = PatternFill("solid", fgColor="1e3a5f")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    bold = Font(bold=True)

    ws_sum["A1"] = f"{brand} — Auditoria Fiscal"
    ws_sum["A1"].font = Font(bold=True, size=14, color="1e3a5f")
    ws_sum["A2"] = f"Período: {month_name}/{year}"
    ws_sum["A3"] = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws_sum.append([])

    summary_rows = [
        ("Score de Risco", f"{period_data.get('period_score', 100)}/100"),
        ("Exposição Total Estimada", _currency(period_data.get("total_exposure", 0))),
        ("Documentos Processados", period_data.get("documents_processed", 0)),
        ("Documentos Críticos", period_data.get("critical_documents", 0)),
        ("Total de Alertas", len(alerts)),
    ]
    for label, value in summary_rows:
        row = ws_sum.max_row + 1
        ws_sum.cell(row=row, column=1, value=label).font = bold
        ws_sum.cell(row=row, column=2, value=value)

    ws_sum.column_dimensions["A"].width = 30
    ws_sum.column_dimensions["B"].width = 22

    # ── Aba 2: Alertas ──
    ws = wb.create_sheet("Alertas")
    headers = [
        "Severidade", "Regra", "Emitente", "Chave Acesso",
        "Data Emissão", "Valor NF (R$)", "Exposição Est. (R$)", "Mensagem",
    ]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    sev_colors = {
        "Crítico": "fef2f2",
        "Alto": "fff7ed",
        "Médio": "fefce8",
        "Baixo": "eff6ff",
    }
    for alert in alerts:
        row = [
            alert["severity_label"],
            alert["rule_id"],
            alert["emitente"],
            alert["chave_acesso"],
            alert["data_emissao"],
            round(alert["valor_nf"], 2),
            round(alert["exposure"], 2),
            alert["message"],
        ]
        ws.append(row)
        fill_color = sev_colors.get(alert["severity_label"], "f8fafc")
        fill = PatternFill("solid", fgColor=fill_color)
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=col_idx).fill = fill

    col_widths = [12, 28, 30, 48, 14, 16, 20, 50]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{ws.max_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
